import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 90, 90)


def set_font(run, size=11, bold=None, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs(paragraph, text, size=11, color=None):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        set_font(run, size=size, bold=bold, color=color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_table(table):
    table.autofit = False
    columns = len(table.columns)
    widths = [9360 // columns] * columns
    widths[-1] += 9360 - sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            shade = "F2F4F7" if row_index == 0 else None
            if shade:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), shade)
                tc_pr.append(shd)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    set_font(run, size=9.5, bold=row_index == 0)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_font(tail, size=9, color=GRAY)


def configure_styles(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10


def parse_table(lines, start):
    if start + 1 >= len(lines) or "|" not in lines[start]:
        return None
    separator = lines[start + 1].strip()
    if not re.match(r"^\|?\s*:?-{3,}", separator):
        return None
    rows = []
    index = start
    while index < len(lines) and "|" in lines[index] and lines[index].strip():
        if index != start + 1:
            rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    return rows, index


def create_document(markdown_path, output_path, title):
    markdown = Path(markdown_path).read_text(encoding="utf-8")
    document = Document()
    configure_styles(document)
    header = document.sections[0].header.paragraphs[0]
    header.text = "技术调研报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=9, color=GRAY)
    add_page_field(document.sections[0].footer.paragraphs[0])

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(6)
    title_paragraph.paragraph_format.space_after = Pt(18)
    title_run = title_paragraph.add_run(title)
    set_font(title_run, size=24, bold=True, color=RGBColor(0, 0, 0))

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        table_data = parse_table(lines, index)
        if table_data:
            rows, index = table_data
            if rows:
                columns = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=columns)
                table.style = "Table Grid"
                for row_index, values in enumerate(rows):
                    for col_index in range(columns):
                        value = values[col_index] if col_index < len(values) else ""
                        paragraph = table.cell(row_index, col_index).paragraphs[0]
                        paragraph.clear()
                        add_runs(paragraph, value, size=9.5)
                configure_table(table)
                document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            value = heading.group(2).strip()
            if value != title:
                paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
                add_runs(paragraph, value, size={1: 16, 2: 13, 3: 12}[len(heading.group(1))])
        elif re.match(r"^[-*]\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+[.)]\s+", "", stripped))
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(8)
            add_runs(paragraph, stripped.lstrip("> "), color=DARK_BLUE)
        else:
            paragraph = document.add_paragraph()
            add_runs(paragraph, stripped)
        index += 1

    document.core_properties.title = title
    document.core_properties.author = "Pi Harness"
    document.core_properties.subject = "Acontext paired document-quality evaluation"
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(json.dumps({"output": str(output), "characters": len(markdown)}, ensure_ascii=False))


def render_manifest(directory):
    root = Path(directory)
    pages = sorted(root.glob("page-*.png"))
    pdfs = sorted(root.glob("*.pdf"))
    print(json.dumps({
        "pageCount": len(pages),
        "pages": [str(path) for path in pages],
        "pdfCount": len(pdfs),
        "pdfs": [str(path) for path in pdfs],
    }, ensure_ascii=False))


def inspect_document(input_path, render_directory):
    path = Path(input_path)
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    rendered_pages = sorted(Path(render_directory).glob("page-*.png"))
    checks = {
        "validDocx": path.exists() and path.stat().st_size > 0,
        "hasTitle": bool(document.core_properties.title),
        "hasHeadings": len(headings) >= 2,
        "substantiveContent": len(text.strip()) >= 800,
        "noPlaceholders": not re.search(r"(?i)\b(TODO|TBD|PLACEHOLDER)\b|<待补充>|\[待补充\]", text),
        "renderedPagesPresent": len(rendered_pages) > 0,
    }
    print(json.dumps({
        "pass": all(checks.values()),
        "checks": checks,
        "paragraphs": len(document.paragraphs),
        "headings": len(headings),
        "tables": len(document.tables),
        "characters": len(text),
        "renderedPages": len(rendered_pages),
        "bytes": path.stat().st_size if path.exists() else 0,
    }, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    create = sub.add_parser("create")
    create.add_argument("--markdown", required=True)
    create.add_argument("--output", required=True)
    create.add_argument("--title", required=True)
    manifest = sub.add_parser("render-manifest")
    manifest.add_argument("--directory", required=True)
    inspect = sub.add_parser("inspect")
    inspect.add_argument("--input", required=True)
    inspect.add_argument("--render-directory", required=True)
    args = parser.parse_args()
    if args.command == "create":
        create_document(args.markdown, args.output, args.title)
    elif args.command == "render-manifest":
        render_manifest(args.directory)
    else:
        inspect_document(args.input, args.render_directory)


if __name__ == "__main__":
    main()
